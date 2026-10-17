"""One workbook with both ledgers and comparable selection evidence."""
from io import BytesIO
from copy import copy
from openpyxl import Workbook,load_workbook
from .paper_export import export_paper_workbook,sheet,TEXT,MONEY,RATE,COUNT,DATE
from .comparison import selection_rules


def export_comparison(state):
    book=Workbook();book.remove(book.active)
    rules=state.get('selection_rules') or selection_rules(state['improved_method'])
    threshold=rules['probability_threshold']
    sheet(book,'비교성적',[
        ('전략','strategy',TEXT,25),('초기금','capital',MONEY,22),('비용 차감 평가액','net_equity',MONEY,24),
        ('수익률','return_pct',RATE,15),('확정 손익','realized_profit',MONEY,22),
        ('최대 낙폭','max_drawdown_pct',RATE,15),('청산 수','trade_count',COUNT,12),
        ('승률','win_rate',RATE,14),('평가 날짜','valuation_date',DATE,16)],
        [{'strategy':'A 기존 전략' if side=='a' else 'B 개선 전략',
          'capital':account['config']['capital'],**account['metrics']} for side,account in state['accounts'].items()])
    sheet(book,'비교조건',[('항목','key',TEXT,27),('내용','value',TEXT,100)],
        [{'key':'시작일','value':state.get('started_day')},
         {'key':'비교 가능 여부','value':str(state['comparison']['comparable'])},
         {'key':'해석','value':state['comparison']['message']},
         {'key':'A','value':'기존 원본 적중률 >=60% + 최신 내부 신호 확률 >=60%'},
         {'key':'B','value':(f"전체 {len((state.get('b_universe') or {}).get('rows',[]))}종목 통합 모델의 비용 차감 수익 확률 상위 {rules['top_k']}종목 (20일 거래대금 하위 {rules['liquidity_cut']:.0%} 제외)" if rules.get('version')==4 else f"비용 차감 수익 확률 >= max({threshold:.0%}, 기본 발생률+{rules['edge_over_base']*100:.0f}%p) + 겹치지 않는 검증 거래 {rules['min_independent_trades']}건 이상·적중률·순수익 우위 + 전체 종목 FDR {rules['fdr_q']:.0%} 다중검정" if rules.get('version')==3 else f'비용 차감 수익 확률 >={threshold:.0%} + 같은 확률 기준의 과거 검증 신호 20개 이상·적중률 우위·순수익 우위')},
         {'key':'B 버전','value':state['improved_method']},
         *[{'key':'B 기준 변경 '+r['changed_at'][:10],'value':r['reason']} for r in state.get('selection_revisions',[])],
         {'key':'공통 체결','value':'다음 종가 매수 / 종가 -3% 손절·+5% 익절 / 각 계좌 10억원 · 종목당 1천만원'},
         *[{'key':side.upper()+' 만기','value':'매수 후 10거래일' if account['config']['holding_mode']=='trading_sessions' else '매수 후 달력 14일'} for side,account in state['accounts'].items()],
         *[{'key':side.upper()+' 분석 대상','value':str(len((account.get('scan') or {}).get('universe',{}).get('rows',[])))+'종목'} for side,account in state['accounts'].items()],
         {'key':'B 구성종목 출처','value':(state.get('b_universe') or {}).get('label','시가총액 상위 코스피50·코스닥50')},
         {'key':'평가','value':'보유분의 예상 매도 비용을 차감한 평가액. 검증 표본은 독립 거래 수가 아님.'}])
    for side,account in state['accounts'].items():
        source=load_workbook(BytesIO(export_paper_workbook(account)))
        for original in source:
            target=book.create_sheet(side.upper()+'_'+original.title)
            for row in original:
                for cell in row:
                    output=target.cell(cell.row,cell.column,cell.value)
                    output.data_type=cell.data_type
                    if cell.has_style:output._style=copy(cell._style)
            for key,dimension in original.column_dimensions.items():target.column_dimensions[key]=copy(dimension)
            target.freeze_panes='A2';target.auto_filter.ref=target.dimensions
    rows=[{**c,'cycle_id':cycle['id']} for cycle in state['accounts']['b'].get('cycles',[])
          for c in cycle['candidates']]
    sheet(book,'B_선정근거',[
        ('차수','cycle_id',COUNT,10),('종목','name',TEXT,25),('코드','code',TEXT,14),
        ('기준일','as_of',DATE,16),('수익 확률','probability',RATE,15),
        ('전략 버전','method',TEXT,30),('선정 확률 기준','threshold',RATE,16),
        ('검증 신호 수','validation.selected_n',COUNT,15),
        ('검증 적중률','validation.selected_hit_rate',RATE,16),('기본 발생률','validation.base_rate',RATE,16),
        ('검증 신호 평균 순수익','validation.selected_mean_return_pct',RATE,24),
        ('전체 검증 평균 순수익','validation.base_mean_return_pct',RATE,24),
        ('학습 끝','training_end',DATE,16),('학습 정답 확정 끝','training_label_end',DATE,20),
        ('결과','status',TEXT,20),('선정 근거','reason',TEXT,65),('오류','error',TEXT,55)],rows)
    output=BytesIO();book.save(output);return output.getvalue()
