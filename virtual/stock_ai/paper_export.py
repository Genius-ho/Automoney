"""Read-only Excel snapshots of the persistent virtual account."""
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from .paper_cycles import cycle_summaries

FILTERS={'all':'전체','candidate':'매수 후보','below':'기준 미달','budget':'1주 예산 부족','error':'분석 제외','duplicate':'중복 제외','cash':'가용 현금 부족'}
CRITERIA={'precision':'원본 검증 적중률 + 상승 신호','probability':'이번 예측의 보정 확률'}
REASONS={'stop_loss':'손절','take_profit':'익절','expiry':'보유 기간 만기'}
TEXT='@';MONEY='#,##0.00';COUNT='#,##0';RATE='0.00%';DATE='yyyy-mm-dd'


def filtered_candidates(rows,result_filter):
    if result_filter not in FILTERS:
        raise ValueError('지원하지 않는 결과 필터입니다.')
    return list(rows) if result_filter=='all' else [r for r in rows if r.get('status')==FILTERS[result_filter]]


def value(row,key,fmt):
    result=row
    for part in key.split('.'):
        result=result.get(part) if isinstance(result,dict) else None
    if result is None:
        return None
    if fmt==DATE and isinstance(result,str):
        return date.fromisoformat(result[:10])
    if fmt==RATE and key.endswith('_pct'):
        return result/100
    if key=='criterion':
        return CRITERIA.get(result,result)
    if key=='reason':
        return REASONS.get(result,result)
    return result


def sheet(book,title,columns,rows):
    ws=book.create_sheet(title)
    ws.append([c[0] for c in columns])
    for row in rows:
        ws.append([value(row,key,fmt) for _,key,fmt,_ in columns])
    for cells in ws:
        for cell in cells:
            if isinstance(cell.value,str):
                cell.value=ILLEGAL_CHARACTERS_RE.sub('',cell.value)
                cell.data_type='s' # Formula-like names remain literal text.
            cell.font=Font(name='맑은 고딕',size=11,color='172B4D')
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            if cell.row>1:
                cell.number_format=columns[cell.column-1][2]
                if cell.row%2==0:
                    cell.fill=PatternFill('solid',fgColor='F1F5F9')
    for cell in ws[1]:
        cell.font=Font(name='맑은 고딕',size=11,bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='17365D')
    ws.row_dimensions[1].height=36
    for index in range(2,ws.max_row+1):
        ws.row_dimensions[index].height=30
    for index,definition in enumerate(columns,1):
        ws.column_dimensions[get_column_letter(index)].width=definition[3]
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    return ws


def export_paper_workbook(state,result_filter='all'):
    candidates=filtered_candidates(state.get('candidates',[]),result_filter)
    book=Workbook();book.remove(book.active)
    config=state['config'];scan=state.get('scan') or {}
    summary=[('총 가상 투자금 (원)',config['capital'],MONEY),('종목당 예산 (원)',state['budget'],MONEY),
        ('현금 (원)',state['cash'],MONEY),('평가액 (원)',state['mark_equity'],MONEY),
        ('누적 실현손익 (원)',sum(t.get('profit',0) for t in state.get('trades',[])),MONEY),
        ('총 수익률',state['return_pct']/100,RATE),('신규 매수','자동 실행' if state['enabled'] else '중지',TEXT),
        ('현재 매수 기준',CRITERIA.get(config['criterion'],config['criterion']),TEXT),
        ('최근 분석 기준일',date.fromisoformat(scan['as_of']) if scan.get('as_of') else None,DATE),
        ('분석 결과 필터',FILTERS[result_filter],TEXT),('분석 결과 행 수',len(candidates),COUNT),
        ('전체 청산 거래 수',len(state.get('trades',[])),COUNT),
        ('다음 차수 재계산일',date.fromisoformat(state['next_scan_day']) if state.get('next_scan_day') else None,DATE),
        ('차수 운용','매주 선정 · 보유/예약 중 종목 중복 제외 · 별도 자금 충전 없음',TEXT),
        ('기록 기준','SQLite 저장 기록. 분석 결과는 선택 필터이며 거래 내역은 전체 이력입니다.',TEXT)]
    overview=sheet(book,'계좌요약',[('항목','label',TEXT,28),('값','value','General',90)],
        [{'label':label,'value':v} for label,v,_ in summary])
    for index,(_,_,fmt) in enumerate(summary,2):
        overview.cell(index,2).number_format=fmt
    sheet(book,'거래내역',[
        ('종목코드','symbol.code',TEXT,14),('종목명','symbol.name',TEXT,26),('시장','symbol.exchange',TEXT,14),
        ('매수일','buy_date',DATE,16),('매도일','sell_date',DATE,16),('수량','quantity',COUNT,12),
        ('매수 기준 종가 (원)','entry_close',MONEY,23),('매수 체결가 (원)','entry_price',MONEY,22),
        ('매도 종가 (원)','sell_close',MONEY,21),('매도 체결가 (원)','sell_price',MONEY,22),
        ('매수 수수료 (원)','buy_fee',MONEY,22),('매도 수수료 (원)','sell_fee',MONEY,22),
        ('매도 거래세 (원)','sell_tax',MONEY,22),('종목 예산 (원)','capital',MONEY,22),
        ('손익 (원)','profit',MONEY,20),('순수익률','net_return_pct',RATE,15),('청산 사유','reason',TEXT,20),
        ('매수 기준','criterion',TEXT,35),('매수 판정값','score',RATE,18),('차수','cycle_id',COUNT,10)],state.get('trades',[]))
    sheet(book,'보유종목',[
        ('종목코드','symbol.code',TEXT,14),('종목명','symbol.name',TEXT,26),('수량','quantity',COUNT,12),
        ('매수일','buy_date',DATE,16),('매수 종가 (원)','entry_close',MONEY,22),('매수 체결가 (원)','entry_price',MONEY,22),
        ('최신 종가 (원)','last_close',MONEY,22),('가격 기준일','last_price_date',DATE,16),('보유 거래일','holding_days',COUNT,16),
        ('차수','cycle_id',COUNT,10),('매도 만기 예정일','expiry_date',DATE,20)],
        state.get('positions',{}).values())
    sheet(book,'매수예약',[
        ('종목코드','symbol.code',TEXT,14),('종목명','symbol.name',TEXT,26),('신호일','signal_date',DATE,16),
        ('매수 예정일','execution_day',DATE,18),('매수 기준','criterion',TEXT,35),('판정값','score',RATE,16),
        ('종목 예산 (원)','budget',MONEY,22),('차수','cycle_id',COUNT,10),('매도 만기 예정일','expiry_date',DATE,20)],
        [{**p,'budget':p['config']['capital']/100} for p in state.get('pending',{}).values()])
    sheet(book,'주문기록',[
        ('종목코드','code',TEXT,14),('종목명','name',TEXT,26),('날짜','day',DATE,16),('주문','type',TEXT,23),
        ('상태','status',TEXT,22),('체결가 (원)','price',MONEY,22),('청산 사유','reason',TEXT,20),('차수','cycle_id',COUNT,10)],state.get('orders',[]))
    score_label='과거 적중률' if scan.get('criterion',config['criterion'])=='precision' else '보정 예측 확률'
    sheet(book,'분석결과',[
        ('시장','exchange',TEXT,14),('시가총액 순위','rank',COUNT,18),('종목코드','code',TEXT,14),('종목명','name',TEXT,26),
        ('기준일','as_of',DATE,16),(score_label,'score',RATE,20),('상승 신호','signal',COUNT,14),
        ('기준 종가 (원)','close',MONEY,22),('결과','status',TEXT,22),('제외 사유','error',TEXT,65),('차수','cycle_id',COUNT,10)],candidates)
    sheet(book,'평가액',[
        ('날짜','date',DATE,16),('평가액 (원)','value',MONEY,24),('현금 (원)','cash',MONEY,24),
        ('보유 종목 수','positions',COUNT,18),('가격 미갱신 종목 수','stale_positions',COUNT,24)],state.get('equity',[]))
    sheet(book,'차수요약',[
        ('차수','id',COUNT,10),('선정일','signal_date',DATE,16),('매수 예정일','execution_day',DATE,18),
        ('매도 만기 예정일','expiry_date',DATE,20),('최초 예약 수','reserved_count',COUNT,16),
        ('중복 제외 수','skipped_duplicate_count',COUNT,16),('현재 예약 수','pending',COUNT,16),
        ('보유 수','held',COUNT,12),('매도 완료 수','sold',COUNT,16),('미체결 종료 수','cancelled',COUNT,18),
        ('확정 손익 (원)','realized_profit',MONEY,24),('상태','status',TEXT,16)],cycle_summaries(state))
    selections=[{**candidate,'cycle_id':cycle['id'],'selection_day':cycle['signal_date'],
                 'execution_day':cycle['execution_day'],'expiry_date':cycle['expiry_date'],
                 'criterion':cycle['criterion']}
        for cycle in state.get('cycles',[]) for candidate in cycle.get('candidates',[])]
    sheet(book,'차수선정기록',[
        ('차수','cycle_id',COUNT,10),('선정일','selection_day',DATE,16),('매수 예정일','execution_day',DATE,18),
        ('매도 만기 예정일','expiry_date',DATE,20),('시장','exchange',TEXT,14),('시가총액 순위','rank',COUNT,18),
        ('종목코드','code',TEXT,14),('종목명','name',TEXT,26),('매수 기준','criterion',TEXT,35),
        ('판정값','score',RATE,16),('상승 신호','signal',COUNT,14),('기준 종가 (원)','close',MONEY,22),
        ('결과','status',TEXT,22),('제외 사유','error',TEXT,65)],selections)
    output=BytesIO();book.save(output)
    return output.getvalue()
