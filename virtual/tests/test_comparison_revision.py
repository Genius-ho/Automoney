from copy import deepcopy
from importlib import import_module
import pytest
from stock_ai.store import Store
from stock_ai.comparison import ComparisonEngine
from test_paper_comparison import Provider, predictions
from test_universe_paper import Universe


def setup_revision(tmp_path):
    store=Store(tmp_path/'ledger.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,
        lambda raw,cfg:predictions(raw,'probability'))
    engine.start('2026-10-07')
    universe=Universe().snapshot();universe['rows']=universe['rows'][:3]
    universe.update(label='구성종목 고정 3종목',expected_count=3)
    before={k:deepcopy(store.get(k)) for k in ['paper_comparison','comparison_paper_a','comparison_paper_b']}
    return engine,provider,universe,before


def apply(engine,provider,universe,**kwargs):
    revise=import_module('stock_ai.comparison_revision').apply_b_selection
    return revise(engine,universe=universe,provider=provider,
        predictor=lambda raw,c:predictions(raw,c),as_of='2026-10-07',
        expected_codes={'000001','000002','000003'},provenance={'source':'test'},**kwargs)


def test_final_b_selection_archives_pending_and_preserves_a(tmp_path):
    engine,provider,universe,before=setup_revision(tmp_path)
    result=apply(engine,provider,universe)
    b=result['accounts']['b'];change=result['selection_revisions'][-1]
    assert engine.store.get('comparison_paper_a')==before['comparison_paper_a']
    assert engine.store.get(change['archive_key'])==before
    assert set(b['pending'])=={'000001','000002','000003'}
    assert b['config']['holding_mode']=='trading_sessions'
    assert all(p['config']['holding_mode']=='trading_sessions' for p in b['pending'].values())
    assert all(p['config']['capital']/100==10_000_000 for p in b['pending'].values())
    assert sum(o['status']=='cancelled' for o in b['orders'])==len(before['comparison_paper_b']['pending'])==20
    assert b['cycles'][0]==before['comparison_paper_b']['cycles'][0]
    assert b['next_scan_day']==before['comparison_paper_b']['next_scan_day']
    assert b['cash']==before['comparison_paper_b']['cash']
    from io import BytesIO
    from openpyxl import load_workbook
    from stock_ai.comparison_export import export_comparison
    book=load_workbook(BytesIO(export_comparison(result)))
    conditions={row[0]:row[1] for row in book['비교조건'].iter_rows(min_row=2,values_only=True)}
    assert conditions['B 만기']=='매수 후 10거래일'
    assert conditions['A 만기']=='매수 후 달력 14일'
    assert book['B_매수예약'].max_row==4
    # Idempotent invocation cannot add a second replacement cycle.
    assert apply(engine,provider,universe)['revision']==result['revision']


def test_b_scope_and_horizon_survive_restart_resume_and_next_scan(tmp_path):
    engine,provider,universe,before=setup_revision(tmp_path)
    apply(engine,provider,universe)
    seen=[]
    def improved(raw,cfg):
        seen.append(cfg.holding_mode);return predictions(raw,'probability')
    fresh=ComparisonEngine(engine.store,provider,Universe(),predictions,improved)
    fresh.stop();fresh.enable('2026-10-07')
    assert fresh.status()['accounts']['a']['config']['holding_mode']=='calendar_days'
    assert fresh.status()['accounts']['b']['config']['holding_mode']=='trading_sessions'
    provider.day=fresh.status()['accounts']['b']['next_scan_day']
    b=fresh.engine('b').scan(provider.day)
    assert len(b['candidates'])==3 and seen==['trading_sessions']*3
    assert len(fresh.engine('a').universe.snapshot()['rows'])==100


@pytest.mark.parametrize('failure',['candidate_mismatch','filled','interrupted'])
def test_revision_failure_does_not_publish_partial_change(tmp_path,failure):
    engine,provider,universe,before=setup_revision(tmp_path)
    if failure=='filled':
        b=engine.store.get('comparison_paper_b');b['orders'][0]['status']='filled'
        engine.store.put('comparison_paper_b',b);before['comparison_paper_b']=deepcopy(b)
    if failure=='candidate_mismatch':provider.missing={'000001'}
    def interrupted(*args):raise RuntimeError('interrupted')
    with pytest.raises((ValueError,RuntimeError)):
        apply(engine,provider,universe,progress=interrupted if failure=='interrupted' else None)
    for key,value in before.items():assert engine.store.get(key)==value
