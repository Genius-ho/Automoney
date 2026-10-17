from copy import deepcopy
from datetime import datetime
from io import BytesIO
import importlib
from openpyxl import load_workbook
import pytest


def sample():
    symbol={'name':'=HYPERLINK("bad")','code':'005930','market':'KR','exchange':'KOSPI'}
    return {'enabled':True,'config':{'capital':1000000000,'criterion':'precision'},'budget':10000000,
        'cash':990000000,'mark_equity':1000500000,'return_pct':.05,'scan':{'as_of':'2026-10-08','criterion':'precision'},
        'positions':{'005930':{'symbol':symbol,'quantity':3,'entry_close':100,'entry_price':100.05,'buy_date':'2026-10-07','last_close':104,'last_price_date':'2026-10-08','holding_days':1}},
        'pending':{'000660':{'symbol':{**symbol,'code':'000660','name':'SK'},'signal_date':'2026-10-08','execution_day':'2026-10-12','score':.7,'criterion':'precision','config':{'capital':1000000000}}},
        'trades':[{'symbol':symbol,'quantity':3,'buy_date':'2026-10-01','sell_date':'2026-10-08','entry_close':100,'entry_price':100.05,'sell_close':96,'sell_price':95.95,'buy_fee':.1,'sell_fee':.2,'sell_tax':.3,'capital':10000000,'profit':-13,'net_return_pct':-.00013,'reason':'stop_loss','criterion':'precision','score':.7} for _ in range(120)],
        'orders':[{'day':'2026-10-12','code':'000660','name':'SK','type':'BUY_NEXT_CLOSE','status':'reserved'}],
        'candidates':[{'exchange':'KOSPI','rank':1,'code':'005930','name':'Samsung','as_of':'2026-10-08','score':.7,'signal':1,'close':100,'status':'매수 후보','eligible':True,'affordable':True},
                      {'exchange':'KOSDAQ','rank':1,'code':'000001','name':'below','as_of':'2026-10-08','score':.5,'close':50,'status':'기준 미달','eligible':False,'affordable':True}],
        'equity':[{'date':'2026-10-08','value':1000500000,'cash':990000000,'positions':1,'stale_positions':0}]}


def export(state,result_filter='all'):
    return importlib.import_module('stock_ai.paper_export').export_paper_workbook(state,result_filter)


def test_excel_contains_all_trades_typed_values_dates_and_safe_text():
    state=sample();before=deepcopy(state)
    workbook=load_workbook(BytesIO(export(state)))
    assert workbook.sheetnames==['계좌요약','거래내역','보유종목','매수예약','주문기록','분석결과','평가액','차수요약','차수선정기록']
    trades=workbook['거래내역']
    assert trades.max_row==121 # all records, not the UI's last100
    assert trades['A2'].value=='005930' and trades['A2'].data_type=='s'
    assert trades['B2'].value==state['trades'][0]['symbol']['name'] and trades['B2'].data_type=='s'
    assert isinstance(trades['D2'].value,datetime)
    headers={cell.value:cell.column for cell in trades[1]}
    assert trades.cell(2,headers['손익 (원)']).value==-13
    assert trades.cell(2,headers['순수익률']).value==pytest.approx(-.0000013)
    assert trades.cell(2,headers['순수익률']).number_format=='0.00%'
    assert trades.cell(2,headers['매도 거래세 (원)']).value==.3
    assert trades.freeze_panes=='A2' and trades.auto_filter.ref.endswith('121')
    assert state==before


def test_excel_filter_only_applies_to_analysis_and_empty_history_is_valid():
    state=sample();state['trades']=[]
    workbook=load_workbook(BytesIO(export(state,'candidate')))
    assert workbook['분석결과'].max_row==2
    assert workbook['분석결과']['C2'].value=='005930'
    assert workbook['거래내역'].max_row==1 and workbook['매수예약'].max_row==2
    with pytest.raises(ValueError):
        export(state,'invalid')


def test_excel_preserves_all_cycle_selections_including_duplicates_and_cycle_ids():
    state=sample()
    for row in state['trades']:row['cycle_id']=1
    state['pending']['000660'].update(cycle_id=2,expiry_date='2026-11-03')
    state['cycles']=[{'id':1,'signal_date':'2026-10-08','execution_day':'2026-10-12',
        'expiry_date':'2026-10-26','criterion':'precision','reserved_count':1,'skipped_duplicate_count':0,
        'candidates':[state['candidates'][0]]},
        {'id':2,'signal_date':'2026-10-19','execution_day':'2026-10-20','expiry_date':'2026-11-03',
         'criterion':'precision','reserved_count':1,'skipped_duplicate_count':1,
         'candidates':[{**state['candidates'][0],'status':'중복 제외','error':'보유 중'},state['candidates'][1]]}]
    workbook=load_workbook(BytesIO(export(state,'candidate')))
    assert workbook['차수요약'].max_row==3 and workbook['차수선정기록'].max_row==4
    selections=workbook['차수선정기록']
    headers={c.value:c.column for c in selections[1]}
    assert selections.cell(3,headers['차수']).value==2
    assert selections.cell(3,headers['결과']).value=='중복 제외'
    assert isinstance(selections.cell(3,headers['매도 만기 예정일']).value,datetime)
    trades=workbook['거래내역'];headers={c.value:c.column for c in trades[1]}
    assert trades.cell(2,headers['차수']).value==1
    assert workbook['분석결과'].max_row==2 # latest filter does not discard archived selections
