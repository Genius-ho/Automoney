from copy import deepcopy
from io import BytesIO
import pandas as pd
import pytest
from stock_ai.store import Store
from stock_ai.universe_paper import UniversePaperEngine,UniverseSettings
from stock_ai.comparison import ComparisonEngine
from stock_ai import strategy_b4
from test_universe_paper import Universe, Quotes


class Provider(Quotes):
    def __init__(self):
        super().__init__()
        self.day='2026-10-07'


def predictions(raw,criterion):
    return {'as_of':str(raw.index[-1].date()),'precision':.8,'probability':.8,'signal':1}


def test_b_v4_reserves_only_the_top_k_without_lowering_a(tmp_path):
    scores=iter(i/200 for i in range(100))
    def prediction(raw,config):
        p=next(scores)
        return {**predictions(raw,config),'precision':p,'probability':p,'entry_allowed':True}
    engine=ComparisonEngine(Store(tmp_path/'db.sqlite'),Provider(),Universe(),lambda raw,c:{**predictions(raw,c),'precision':.5},prediction)
    state=engine.start('2026-10-07')
    assert not state['accounts']['a']['pending']
    b=state['accounts']['b']
    # codes 000001..000100 get scores 0.000..0.495 in order -> the last 20 are the best
    assert sorted(b['pending'])==[str(i).zfill(6) for i in range(81,101)]
    assert b['scan']['selection_threshold']==0 and state['selection_rules']['top_k']==20
    ranked={c['code']:c for c in b['candidates']}
    assert ranked['000100']['pool_rank']==1 and ranked['000080']['status']=='기준 미달'


def test_b_revision_preserves_a_and_archives_empty_old_b_atomically(tmp_path):
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,
        lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':False})
    engine.start('2026-10-07')
    meta=store.get('paper_comparison');meta['improved_method']='trade-outcome-oof-v1';store.put('paper_comparison',meta)
    before={k:deepcopy(store.get(k)) for k in ['paper_comparison','comparison_paper_a','comparison_paper_b']}
    engine.improved=lambda raw,cfg:{**predictions(raw,cfg),'probability':.55,'entry_allowed':True}
    result=engine.revise_untraded_b()
    assert store.get('comparison_paper_a')==before['comparison_paper_a']
    assert len(result['accounts']['b']['pending'])==100
    assert result['accounts']['b']['cycles'][0]==before['comparison_paper_b']['cycles'][0]
    change=result['selection_revisions'][-1]
    assert change['from_method']=='trade-outcome-oof-v1'
    assert store.get(change['archive_key'])==before
    assert result['selection_rules']['probability_threshold']==.52
    assert result['accounts']['b']['pending']['000001']['execution_day']=='2026-10-08'


def test_failed_b_revision_keeps_previous_method_and_accounts(tmp_path):
    store=Store(tmp_path/'db.sqlite')
    engine=ComparisonEngine(store,Provider(),Universe(),predictions,
        lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':False})
    engine.start('2026-10-07')
    meta=store.get('paper_comparison');meta['improved_method']='trade-outcome-oof-v1';store.put('paper_comparison',meta)
    before={k:deepcopy(store.get(k)) for k in ['paper_comparison','comparison_paper_a','comparison_paper_b']}
    def failure(percent,message):raise RuntimeError('interrupted')
    with pytest.raises(RuntimeError,match='interrupted'):engine.revise_untraded_b(progress=failure)
    for key,value in before.items():assert store.get(key)==value


def test_b_revision_refuses_to_mix_completed_trading_periods(tmp_path):
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,
        lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':False})
    engine.start('2026-10-07');provider.day='2026-10-08';engine.daily('2026-10-08')
    meta=store.get('paper_comparison');meta['improved_method']='trade-outcome-oof-v1';store.put('paper_comparison',meta)
    with pytest.raises(ValueError,match='체결'):engine.revise_untraded_b()
    assert store.get('paper_comparison')==meta


def test_b_revision_rejects_changed_prices_without_publishing(tmp_path):
    store=Store(tmp_path/'db.sqlite')
    engine=ComparisonEngine(store,Provider(),Universe(),predictions,
        lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':False,'data_hash':'first-prices'})
    engine.start('2026-10-07')
    meta=store.get('paper_comparison');meta['improved_method']='trade-outcome-oof-v1';store.put('paper_comparison',meta)
    before={key:deepcopy(store.get(key)) for key in ['paper_comparison','comparison_paper_a','comparison_paper_b']}
    engine.improved=lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':True,'data_hash':'revised-prices'}
    with pytest.raises(ValueError,match='주가 자료'):engine.revise_untraded_b()
    for key,value in before.items():assert store.get(key)==value


def test_namespaced_accounts_cannot_modify_original(tmp_path):
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    original=UniversePaperEngine(store,provider,Universe(),predictions)
    original.configure(UniverseSettings().model_dump(),False)
    before=deepcopy(store.get('universe_paper'))
    other=UniversePaperEngine(store,provider,Universe(),predictions,state_key='other')
    other.configure(UniverseSettings().model_dump(),True)
    other.scan('2026-10-07')
    assert len(other.status()['pending'])==100
    assert store.get('universe_paper')==before


def test_comparison_uses_same_universe_and_prices_but_independent_selection(tmp_path):
    class CountingProvider(Provider):
        def __init__(self):
            super().__init__();self.calls=[]
        def original_prices(self,code,market):
            self.calls.append(code)
            return super().original_prices(code,market)
    store=Store(tmp_path/'db.sqlite');provider=CountingProvider()
    before=store.get('universe_paper')
    seen={'a':[],'b':[]}
    def a(raw,criterion):
        seen['a'].append(raw.to_json());return predictions(raw,criterion)
    def b(raw,config):
        seen['b'].append(raw.to_json());return {**predictions(raw,'probability'),'entry_allowed':False,'reason':'검증 우위 미달'}
    engine=ComparisonEngine(store,provider,Universe(),a,b)
    engine.start('2026-10-07')
    state=engine.status()
    assert len(provider.calls)==100 and seen['a']==seen['b']
    assert len(state['accounts']['a']['pending'])==100
    assert state['accounts']['b']['pending']=={}
    assert state['accounts']['a']['cash']==state['accounts']['b']['cash']==1e9
    assert state['started_day']=='2026-10-07'
    assert store.get('universe_paper')==before
    restored=ComparisonEngine(store,provider,Universe(),a,b).status()
    assert restored['accounts']['a']['pending']==state['accounts']['a']['pending']
    engine.stop()
    assert not engine.status()['enabled']
    assert not engine.status()['accounts']['a']['pending']


def test_comparison_settles_same_closes_and_reports_returns_without_reset(tmp_path,monkeypatch):
    # same 100 selections on both sides isolates settlement from B's top-K cut
    monkeypatch.setattr(strategy_b4,'TOP_K',100)
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,lambda raw,cfg:predictions(raw,'probability'))
    engine.start('2026-10-07')
    provider.day='2026-10-08'
    engine.daily('2026-10-08')
    state=engine.status()
    assert len(state['accounts']['a']['positions'])==100
    assert state['accounts']['a']['cash']==state['accounts']['b']['cash']
    assert state['comparison']['comparable']
    assert state['comparison']['return_difference_pct']==0
    old=deepcopy(state['accounts']['a']['positions'])
    engine.start('2026-10-08')
    assert engine.status()['accounts']['a']['positions']==old
    assert len(engine.status()['accounts']['a']['cycles'])==1


def test_missing_quotes_do_not_produce_a_winner(tmp_path):
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,lambda raw,cfg:predictions(raw,'probability'))
    engine.start('2026-10-07');provider.day='2026-10-08';engine.daily('2026-10-08')
    state=store.get('comparison_paper_a')
    next(iter(state['positions'].values()))['last_price_date']='2026-10-07'
    store.put('comparison_paper_a',state)
    assert not engine.status()['comparison']['comparable']


def test_failed_settlement_replays_missed_stop_before_rebound(tmp_path,monkeypatch):
    # same 100 selections on both sides isolates settlement from B's top-K cut
    monkeypatch.setattr(strategy_b4,'TOP_K',100)
    class HistoryProvider(Provider):
        def original_prices(self,code,market='KR'):
            if code in self.missing:raise ValueError('unavailable')
            dates=pd.to_datetime(['2026-10-07','2026-10-08','2026-10-12','2026-10-13'])
            prices=[100,100,96,100]
            raw=pd.DataFrame({k:prices for k in ['Open','High','Low','Close']},index=dates)
            raw['Volume']=100;raw['Change']=0
            return raw.loc[:self.day],{'source':'fixture'}
    store=Store(tmp_path/'db.sqlite');provider=HistoryProvider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,lambda raw,cfg:predictions(raw,'probability'))
    engine.start('2026-10-07');provider.day='2026-10-08';engine.daily('2026-10-08')
    before={side:deepcopy(store.get('comparison_paper_'+side)) for side in ['a','b']}
    provider.day='2026-10-13';provider.missing={'000001'}
    failed=engine.daily('2026-10-13')
    assert failed['last_settled_day']=='2026-10-08'
    assert failed['incomplete_day']=='2026-10-12'
    assert not failed['comparison']['comparable']
    for side in before:assert store.get('comparison_paper_'+side)==before[side]
    provider.missing.clear()
    recovered=engine.daily('2026-10-13')
    assert recovered['last_settled_day']=='2026-10-13'
    assert not recovered.get('incomplete_day')
    assert recovered['comparison']['comparable']
    for state in recovered['accounts'].values():
        assert not state['positions']
        assert len(state['trades'])==100
        assert {t['sell_date'] for t in state['trades']}=={'2026-10-12'}
        assert {t['reason'] for t in state['trades']}=={'stop_loss'}


def test_fully_stopped_cohort_is_refilled_before_the_weekly_scan(tmp_path,monkeypatch):
    monkeypatch.setattr(strategy_b4,'TOP_K',100)
    class DropProvider(Provider):
        def original_prices(self,code,market='KR'):
            dates=pd.to_datetime(['2026-10-07','2026-10-08','2026-10-12','2026-10-13'])
            # 000001-000050 fall -4% after entry (stop), the rest stay flat
            drop=int(code)<=50
            prices=[100,100,96 if drop else 100,96 if drop else 100]
            raw=pd.DataFrame({k:prices for k in ['Open','High','Low','Close']},index=dates)
            raw['Volume']=100;raw['Change']=0
            return raw.loc[:self.day],{'source':'fixture'}
    store=Store(tmp_path/'db.sqlite');provider=DropProvider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,lambda raw,cfg:predictions(raw,'probability'))
    engine.start('2026-10-07');provider.day='2026-10-08';engine.daily('2026-10-08')
    for side in ('a','b'):
        state=store.get('comparison_paper_'+side)
        assert state['next_scan_day']>'2026-10-12' and len(state['positions'])==100
        # make the flat half expire on the same day so the whole cohort is closed on 10-12
        for code in [c for c in state['positions'] if int(c)>50]:
            state['positions'][code]['expiry_date']='2026-10-12'
        store.put('comparison_paper_'+side,state)
    provider.day='2026-10-12';result=engine.daily('2026-10-12')
    for side in ('a','b'):
        state=result['accounts'][side]
        # cohort closed before the weekly slot -> refilled on the same settlement day
        assert state['cycles'][-1].get('refill') and state['cycles'][-1]['signal_date']=='2026-10-12'
        stopped={t['symbol']['code'] for t in state['trades'] if t['reason']=='stop_loss'}
        # no cooldown: stopped stocks are selected again whenever the rules pick them
        assert len(stopped)==50 and stopped<=set(state['pending'])
    # the refilled cohort is open: no further early rescan and the weekly slot is unchanged
    assert not any(engine.engine(s).refill_due('2026-10-13') for s in ('a','b'))
    assert all(result['accounts'][s]['next_scan_day']==store.get('comparison_paper_'+s)['next_scan_day'] for s in ('a','b'))


def test_stopped_positions_are_reentered_one_for_one_before_the_weekly_scan(tmp_path,monkeypatch):
    monkeypatch.setattr(strategy_b4,'TOP_K',100)
    class DropProvider(Provider):
        def original_prices(self,code,market='KR'):
            dates=pd.to_datetime(['2026-10-07','2026-10-08','2026-10-12','2026-10-13'])
            # 000001-000030 stop out on 10-12, the other 70 stay flat and keep holding
            drop=int(code)<=30
            prices=[100,100,96 if drop else 100,96 if drop else 100]
            raw=pd.DataFrame({k:prices for k in ['Open','High','Low','Close']},index=dates)
            raw['Volume']=100;raw['Change']=0
            return raw.loc[:self.day],{'source':'fixture'}
    store=Store(tmp_path/'db.sqlite');provider=DropProvider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,lambda raw,cfg:predictions(raw,'probability'))
    engine.start('2026-10-07');provider.day='2026-10-08';engine.daily('2026-10-08')
    provider.day='2026-10-12';result=engine.daily('2026-10-12')
    for side in ('a','b'):
        state=result['accounts'][side]
        cycle=state['cycles'][-1]
        # only the 30 freed slots are re-reserved, never the 70 stocks still held. The account is fully
        # invested, so the -4% losses leave cash for 28 of the 30 budgets: the last two wait for the weekly scan.
        assert cycle.get('refill') and cycle['reentry_slots']==30 and cycle['reserved_count']==28
        assert len(state['pending'])==28 and not set(state['pending'])&set(state['positions'])
        assert sum(c['status']=='가용 현금 부족' for c in state['candidates'])==2
        assert len(state['positions'])==70


def test_stop_at_final_publish_does_not_open_nested_transaction(tmp_path):
    from contextlib import contextmanager
    class StrictStore(Store):
        active=False
        @contextmanager
        def transaction(self):
            assert not self.active,'nested transaction'
            self.active=True
            try:
                with super().transaction() as con:yield con
            finally:self.active=False
    engine=ComparisonEngine(StrictStore(tmp_path/'db.sqlite'),Provider(),Universe(),predictions,
        lambda raw,cfg:predictions(raw,'probability'))
    engine.start('2026-10-07')
    def stop_at_end(percent,message):
        if percent==98:engine.stop()
    state=engine.daily('2026-10-07',stop_at_end)
    assert not state['enabled']
    assert all(not account['pending'] for account in state['accounts'].values())


def test_rejected_start_preserves_existing_enabled_accounts(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from stock_ai.web import create_app
    provider=Provider()
    app=create_app(tmp_path,provider,scheduler=False)
    engine=app.state.comparison
    engine.universe=Universe();engine.baseline=predictions
    engine.improved=lambda raw,cfg:predictions(raw,'probability')
    engine.start('2026-10-07')
    before={key:deepcopy(app.state.store.get(key)) for key in
        ['paper_comparison','comparison_paper_a','comparison_paper_b']}
    def reject(*args,**kwargs):raise ValueError('queue full')
    monkeypatch.setattr(app.state.jobs,'submit',reject)
    with TestClient(app) as client:
        assert client.post('/api/comparison-paper/start').status_code==409
    for key in before:assert app.state.store.get(key)==before[key]


def test_stop_during_analysis_discards_both_staged_accounts(tmp_path):
    engine=ComparisonEngine(Store(tmp_path/'db.sqlite'),Provider(),Universe(),predictions,
        lambda raw,cfg:predictions(raw,'probability'))
    stopped=[]
    def progress(percent,message):
        if message.startswith('A · 1/100') and not stopped:
            stopped.append(True);engine.stop()
    result=engine.start('2026-10-07',progress)
    assert not result['enabled']
    for state in result['accounts'].values():
        assert not state['pending'] and not state['cycles']


def test_failure_between_strategies_cannot_publish_one_sided_orders(tmp_path):
    engine=ComparisonEngine(Store(tmp_path/'db.sqlite'),Provider(),Universe(),predictions,
        lambda raw,cfg:predictions(raw,'probability'))
    def fail(percent,message):
        if message.startswith('B · 1/100'):raise RuntimeError('worker failed')
    with pytest.raises(RuntimeError,match='worker failed'):
        engine.start('2026-10-07',fail)
    for state in engine.status()['accounts'].values():
        assert not state['pending'] and not state['cycles']


def test_comparison_api_and_excel_preserve_two_independent_ledgers(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from stock_ai.web import create_app
    from test_web import wait_for_job
    provider=Provider();provider.day='2026-10-08'
    monkeypatch.setattr(UniversePaperEngine,'completed_day',lambda self:'2026-10-08')
    app=create_app(tmp_path,provider,scheduler=False)
    engine=app.state.comparison
    engine.universe=Universe();engine.baseline=predictions
    engine.improved=lambda raw,cfg:{**predictions(raw,'probability'),'entry_allowed':False}
    with TestClient(app) as client:
        before=app.state.store.get('universe_paper')
        assert client.get('/api/comparison-paper').json()['started_day'] is None
        job=wait_for_job(client,client.post('/api/comparison-paper/start').json()['id'])
        assert job['status']=='done',job['error']
        state=client.get('/api/comparison-paper').json()
        assert len(state['accounts']['a']['pending'])==100
        assert not state['accounts']['b']['pending']
        response=client.get('/api/comparison-paper/export.xlsx')
        book=load_workbook(BytesIO(response.content))
        assert book['A_매수예약'].max_row==101
        assert book['B_매수예약'].max_row==1
        assert book['B_선정근거'].max_row==101
        assert client.post('/api/comparison-paper/stop').json()['enabled'] is False
        assert app.state.store.get('universe_paper')==before


def test_stored_v2_keeps_settling_and_scanning_a_but_waits_to_select_b(tmp_path):
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,
        lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':True})
    engine.start('2026-10-07')
    meta=store.get('paper_comparison');meta['improved_method']='trade-outcome-oof-v2-p52';store.put('paper_comparison',meta)
    b_before=deepcopy(store.get('comparison_paper_b'))
    provider.day='2026-10-08';state=engine.daily('2026-10-08')
    assert state['last_settled_day']=='2026-10-08' and 'B 저장 버전' in state['b_selection_blocked']
    # yesterday's B reservations were still filled at the 10-08 close
    assert len(state['accounts']['b']['positions'])==len(b_before['pending'])
    assert state['accounts']['b']['cycles']==b_before['cycles']


@pytest.mark.parametrize('previous',['trade-outcome-oof-v2-p52','trade-outcome-wf-v3-edge-fdr'])
def test_upgrade_to_v4_archives_and_keeps_existing_b_reservations(tmp_path,previous):
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,
        lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':True})
    engine.start('2026-10-07')
    meta=store.get('paper_comparison');meta['improved_method']=previous;store.put('paper_comparison',meta)
    before={k:deepcopy(store.get(k)) for k in ['paper_comparison','comparison_paper_a','comparison_paper_b']}
    result=engine.upgrade_b_method('test',rescan_now=True)
    change=result['selection_revisions'][-1]
    assert result['improved_method']=='pooled-rank-top-k-v4' and result['selection_rules']['top_k']==20
    assert change['from_method']==previous and store.get(change['archive_key'])==before
    assert store.get('comparison_paper_a')==before['comparison_paper_a']
    # reservations exist, so the weekly slot is kept even with rescan_now
    assert store.get('comparison_paper_b')==before['comparison_paper_b']
    assert change['kept_pending']==sorted(before['comparison_paper_b']['pending'])
    assert engine.upgrade_b_method('again')['revision']==result['revision']
    meta=store.get('paper_comparison');meta['improved_method']='trade-outcome-oof-v1';store.put('paper_comparison',meta)
    with pytest.raises(ValueError,match='v2'):engine.upgrade_b_method('x')


def test_upgrade_with_rescan_lets_an_empty_b_select_at_the_next_settlement(tmp_path):
    store=Store(tmp_path/'db.sqlite');provider=Provider()
    engine=ComparisonEngine(store,provider,Universe(),predictions,
        lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':False})
    engine.start('2026-10-07')
    # the stored 2026-10-13 state: a v3 B that selected nothing and waits a week
    meta=store.get('paper_comparison');meta['improved_method']='trade-outcome-wf-v3-edge-fdr';store.put('paper_comparison',meta)
    assert not store.get('comparison_paper_b')['pending'] and store.get('comparison_paper_b')['next_scan_day']>'2026-10-08'
    engine.improved=lambda raw,cfg:{**predictions(raw,cfg),'entry_allowed':True}
    engine.upgrade_b_method('v3 selected nothing',rescan_now=True)
    a_before=deepcopy(store.get('comparison_paper_a'))
    provider.day='2026-10-08';state=engine.daily('2026-10-08')
    assert len(state['accounts']['b']['pending'])==20
    assert state['accounts']['b']['pending']['000001']['execution_day']=='2026-10-12'
    # A keeps its own weekly slot: no new A cycle
    assert len(state['accounts']['a']['cycles'])==len(a_before['cycles'])
