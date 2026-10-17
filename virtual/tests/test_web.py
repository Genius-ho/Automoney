from fastapi.testclient import TestClient
from stock_ai.web import create_app
from stock_ai.config import Settings
from test_paper_service import FakeProvider
import time
import pytest


def wait_for_job(client, identifier, timeout=120):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get('/api/jobs/' + identifier).json()
        if job['status'] in ('done', 'failed'):
            return job
        time.sleep(.05)
    pytest.fail('Background job did not finish before the deadline')


def test_universe_paper_endpoints_preserve_existing_account(tmp_path):
    with TestClient(create_app(tmp_path,FakeProvider(),scheduler=False)) as client:
        old=client.get('/api/paper').json()
        initial=client.get('/api/universe-paper')
        assert initial.status_code==200
        assert initial.json()['enabled'] is False and initial.json()['budget']==10000000
        response=client.post('/api/universe-paper/configure',json={'config':{'capital':200000000,'criterion':'precision'},'enabled':False})
        assert response.status_code==200 and response.json()['budget']==2000000
        assert client.get('/api/paper').json()==old
        assert client.post('/api/universe-paper/configure',json={'config':{'capital':0},'enabled':True}).status_code==422
        assert client.post('/api/universe-paper/configure',json={'config':{'criterion':'unknown'},'enabled':True}).status_code==422


def test_activity_api_returns_planned_dates_and_closing_quote_without_changing_reservations(tmp_path):
    app=create_app(tmp_path,FakeProvider(),scheduler=False)
    state=app.state.universe_paper.status()
    state['pending']={'005930':{'symbol':{'code':'005930','name':'Samsung','market':'KR'},
        'execution_day':'2026-10-12','signal_date':'2026-10-08'}}
    state['candidates']=[{'code':'005930','close':70000,'as_of':'2026-10-08'}]
    app.state.store.put('universe_paper',state)
    with TestClient(app) as client:
        before=app.state.store.get('universe_paper')
        response=client.get('/api/universe-paper')
        assert response.status_code==200
        row=response.json()['activity'][0]
        assert row['buy_date']=='2026-10-12' and row['sell_date']=='2026-10-26'
        assert row['current_price']==70000 and row['average_price'] is None
        assert row['profit'] is None
        assert app.state.store.get('universe_paper')==before


def test_paper_excel_download_preserves_database_and_filters_only_analysis(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    from test_paper_export import sample
    app=create_app(tmp_path,FakeProvider(),scheduler=False)
    state=sample()
    app.state.store.put('universe_paper',state)
    with TestClient(app) as client:
        before=app.state.store.get('universe_paper')
        response=client.get('/api/universe-paper/export.xlsx?result_filter=candidate')
        assert response.status_code==200
        assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        assert 'attachment' in response.headers['content-disposition']
        book=load_workbook(BytesIO(response.content))
        assert book['거래내역'].max_row==121 and book['분석결과'].max_row==2
        assert app.state.store.get('universe_paper')==before
        assert client.get('/api/universe-paper/export.xlsx?result_filter=unknown').status_code==422


def test_web_settings_mapping_and_path_validation(tmp_path):
    with TestClient(create_app(tmp_path,FakeProvider(),scheduler=False)) as client:
        assert client.get('/').status_code==200
        settings=client.get('/api/settings').json()
        assert settings['horizon']==10
        assert client.put('/api/settings',json={**settings,'horizon':0}).status_code==422
        assert client.put('/api/settings',json={**settings,'horizon':12}).status_code==200
        assert client.get('/api/symbols',params={'q':'삼성에스디에스'}).json()['candidates'][0]['code']=='018260'
        assert client.get('/api/symbols',params={'q':'삼성저낮'}).json()['candidates'][0]['match_type']=='fuzzy'
        etf=client.get('/api/symbols',params={'q':'비트코인 ETF','market':'US'}).json()['candidates'][0]
        assert etf['code']=='IBIT' and etf['asset_type']=='ETF'
        assert client.get('/api/symbols',params={'q':'GLD','market':'KR'}).json()['candidates']==[]
        assert client.get('/api/symbols',params={'q':'GLD','market':'bad'}).status_code==422
        assert client.get('/api/reports/not-a-uuid/result.json').status_code in [404,422]
        bad={'symbol':{'code':'AAPL','name':'Apple','market':'US'},'settings':settings,'enabled':True}
        assert client.post('/api/paper/configure',json=bad).status_code==400
    with TestClient(create_app(tmp_path,FakeProvider(),scheduler=False)) as client:
        assert client.get('/api/settings').json()['horizon']==12


def test_analysis_job_validates_symbol_and_csv_errors_visible(tmp_path):
    with TestClient(create_app(tmp_path,FakeProvider(),scheduler=False)) as client:
        assert client.post('/api/analyze',json={'symbol':{'code':'../../bad','name':'bad','market':'KR'}}).status_code==422
        response=client.post('/api/analyze',json={'symbol':{'code':'005930','name':'삼성전자','market':'KR'},'csv':'not,prices\n1,2'})
        assert response.status_code==202
        import time
        for _ in range(100):
            job=client.get('/api/jobs/'+response.json()['id']).json()
            if job['status']=='failed':break
            time.sleep(.01)
        assert job['status']=='failed'
        assert '열' in job['error']


def test_csv_analysis_result_reports_and_job_survive_restart(tmp_path, prices):
    with TestClient(create_app(tmp_path, FakeProvider(), scheduler=False)) as client:
        response = client.post('/api/analyze', json={
            'symbol': {'code': '005930', 'name': '삼성전자', 'market': 'KR'},
            'csv': prices.to_csv()})
        assert response.status_code == 202
        job = wait_for_job(client, response.json()['id'])
        assert job['status'] == 'done', job['error']
        result = job['result']
        assert result['feature_count'] == 11
        assert result['analysis_method'] == 'original-close-v1'
        assert 0 <= result['mean_precision'] <= 1
        assert result['latest_signal'] in (0,1)
        assert 'latest_probability' not in result
        assert result['data_source']['source'] == '사용자 CSV'
        assert set(result['simulations']) == {'touch', 'buffer', 'stress'}
        report_id = result['report_id']
        for filename in ('result.json', 'report.md', 'predictions.csv', 'trades_buffer.csv', 'equity_buffer.csv'):
            assert client.get(f'/api/reports/{report_id}/{filename}').status_code == 200
        assert client.get(f'/api/reports/{report_id}/secret.txt').status_code == 404
        assert client.get('/api/latest').json()['data_hash'] == result['data_hash']
    with TestClient(create_app(tmp_path, FakeProvider(), scheduler=False)) as client:
        assert client.get('/api/jobs/' + job['id']).json()['status'] == 'done'
        assert client.get('/api/latest').json()['report_id'] == report_id


@pytest.mark.parametrize('path,payload', [
    ('/api/batch', {}), ('/api/symbols/refresh', None),
    ('/api/portfolio', {'symbols': [
        {'code': '005930', 'name': '삼성전자', 'market': 'KR'},
        {'code': '009150', 'name': '삼성전기', 'market': 'KR'}]}),
    ('/api/paper/prepare', None),
])
def test_full_queue_returns_conflict_instead_of_server_error(tmp_path, path, payload):
    from threading import Event
    release = Event()
    app = create_app(tmp_path, FakeProvider(), scheduler=False)
    with TestClient(app, raise_server_exceptions=False) as client:
        app.state.paper.configure({'code': '005930', 'name': '삼성전자', 'market': 'KR'},
                                  Settings(), True)
        try:
            for _ in range(3):
                app.state.jobs.submit('blocked', lambda progress: release.wait(10))
            response = client.post(path, json=payload)
            assert response.status_code == 409
            assert isinstance(response.json()['detail'], str)
        finally:
            release.set()


def test_both_kelly_formulas_are_available(tmp_path):
    with TestClient(create_app(tmp_path, FakeProvider(), scheduler=False)) as client:
        first = client.post('/api/kelly', json={'probability': .6, 'payoff': 2}).json()
        second = client.post('/api/kelly', json={'probability': .6, 'payoff': .2, 'loss': .1}).json()
        assert first['raw_fraction'] == pytest.approx(.4)
        assert second['raw_fraction'] == pytest.approx(4)
        assert second['long_only_fraction'] == 1


def test_portfolio_budget_result_survives_restart_and_preserves_virtual_account(tmp_path):
    from test_portfolio_service import Quotes,SYMBOLS
    app=create_app(tmp_path,Quotes(),scheduler=False)
    app.state.universe_paper.configure({'capital':1000000000},False)
    before=app.state.store.get('universe_paper')
    with TestClient(app) as client:
        assert client.get('/api/portfolio/latest').json() is None
        request=client.post('/api/portfolio',json={'symbols':SYMBOLS,'years':5,'capital':1000000})
        assert request.status_code==202
        job=wait_for_job(client,request.json()['id'])
        assert job['status']=='done',job.get('error')
        assert job['request']['symbols']==SYMBOLS
        assert job['request']['capital']==1000000
        result=client.get('/api/portfolio/latest').json()
        assert result==job['result'] and len(result['allocations'])==2
        assert result['capital']==1000000 and result['currency']=='KRW'
        assert result['portfolio_strategy']=='half-kelly-v1'
        assert result['constraints']['min_weight_pct']==5
        assert result['reserve_cash']>=500000
        assert app.state.store.get('universe_paper')==before
        assert client.post('/api/portfolio',json={'symbols':[SYMBOLS[0],SYMBOLS[0]]}).status_code==422
        assert client.post('/api/portfolio',json={'symbols':SYMBOLS,'currency':'invalid'}).status_code==422
        assert client.post('/api/portfolio',json={'symbols':SYMBOLS,'capital':0}).status_code==422
        assert client.post('/api/portfolio',json={'symbols':SYMBOLS+[{'code':'068270','name':'Cell','market':'KR'}],
                           'max_weight_pct':30}).status_code==422
        assert client.post('/api/portfolio',json={'symbols':[
            {'code':'AAPL','name':'Apple','market':'US'},
            {'code':'aapl','name':'Apple','market':'US'}]}).status_code==422
    with TestClient(create_app(tmp_path,Quotes(),scheduler=False)) as client:
        assert client.get('/api/portfolio/latest').json()==result


def test_portfolio_accepts_mixed_stocks_and_saves_basis_currency(tmp_path):
    import pandas as pd
    from test_portfolio_service import Quotes,SYMBOLS
    class MixedQuotes(Quotes):
        def exchange_rates(self,years):
            return pd.DataFrame({'Close':1300},index=pd.bdate_range('2025-01-02',periods=100)),{}
    with TestClient(create_app(tmp_path,MixedQuotes(),scheduler=False)) as client:
        symbols=[SYMBOLS[0],{'code':'GLD','name':'Gold','market':'US','name_en':'SPDR Gold Shares','industry':'금','asset_type':'ETF'}]
        response=client.post('/api/portfolio',json={'symbols':symbols,'currency':'KRW','capital':10000000})
        assert response.status_code==202
        job=wait_for_job(client,response.json()['id'])
        assert job['status']=='done',job.get('error')
        assert job['result']['currency']=='KRW' and job['request']['currency']=='KRW'
        assert job['result']['exchange_rate']['rate']==1300
        assert job['request']['symbols'][1]['name_en']=='SPDR Gold Shares'
        assert job['result']['allocations'][1]['symbol']['industry']=='금'


def test_saved_foreign_analysis_has_korean_display_metadata_without_changing_stored_record(tmp_path):
    app=create_app(tmp_path,FakeProvider(),scheduler=False)
    original={'symbol':{'code':'MU','name':'MU','market':'US'}}
    app.state.store.put('latest_analysis',original)
    with TestClient(app) as client:
        latest=client.get('/api/latest').json()
        assert '마이크론' in latest['symbol']['name'] and 'Micron' in latest['symbol']['name_en']
        assert app.state.store.get('latest_analysis')==original
